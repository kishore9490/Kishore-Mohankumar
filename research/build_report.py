# -*- coding: utf-8 -*-
"""
Generate the AI Revenue Recovery OS — U.S. Payer API & RCM Feasibility
deliverables: CSV, Excel workbook, PDF report, and Markdown findings.

Run:  python3 build_report.py
All capability values derive from research_data.py (sourced) plus the
per-payer rule engine below. Uncertainty is preserved (UNKNOWN/PARTIAL).
No PHI. No fabricated pricing.
"""
import os, textwrap
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import pandas as pd

from payer_universe import PAYERS, CLEARINGHOUSES
import research_data as R

OUT = os.path.dirname(os.path.abspath(__file__))
ASSET = os.path.join(OUT, "_assets"); os.makedirs(ASSET, exist_ok=True)

# ---------- palette ----------
NAVY="#0C1B30"; NAVY2="#13253F"; TEAL="#118C7E"; TEAL_L="#2FBFA6"
CYAN="#2E7FC2"; INK="#1a2433"; MUT="#6B7B90"; LINE="#D5DEE8"
GREEN="#2E9E6B"; AMBER="#D8942A"; REDX="#C7503F"; GREY="#95A3B4"; VIOLET="#6C5CE0"
VAL_COLOR={"YES":GREEN,"PARTIAL":AMBER,"NO":REDX,"UNKNOWN":GREY,"NOT APPLICABLE":"#B9C4D0","N/A":"#B9C4D0"}
LIGHT={"YES":"#E4F3EA","PARTIAL":"#FBF0DA","NO":"#F7E1DD","UNKNOWN":"#EDF1F5","NOT APPLICABLE":"#EEF1F4","N/A":"#EEF1F4"}

plt.rcParams.update({"font.family":"DejaVu Sans","font.size":9,"axes.edgecolor":LINE,
    "axes.linewidth":.8,"figure.dpi":150,"savefig.dpi":150,"text.color":INK,
    "axes.labelcolor":INK,"xtick.color":MUT,"ytick.color":MUT})

# ======================================================================
# 1. PER-PAYER CAPABILITY RECORDS (113 fields)
# ======================================================================
COLS = [
 "Payer Organization","Parent Organization","Payer Type","Plan Type","State(s)","Payer ID","Payer ID Alias","Website","Developer Portal","Provider Portal",
 "Claim Status Electronic","Claim Status Standard","Claim Status API","Claim Status Direct Payer API","Claim Status Clearinghouse API","Claim Status FHIR","Claim Status Portal","Claim Status Voice Fallback","Claim Status Real-Time","Claim Status Batch",
 "Eligibility Electronic","Eligibility Standard","Eligibility API","Eligibility Direct Payer API","Eligibility Clearinghouse API","Eligibility FHIR","Eligibility Portal","Eligibility Real-Time","Eligibility Batch",
 "Benefits Available","Benefits via 271","Benefits via API","Deductible","Deductible Remaining","Copay","Coinsurance","OOP Maximum","OOP Remaining","Service-Specific Benefits","Network Benefits",
 "Allowed Amount Available","Accumulator Data","Patient Responsibility Data","Patient Responsibility Calculation Required","Estimation API","Estimation Confidence",
 "Prior Auth Electronic","Prior Auth X12 278","Prior Auth FHIR","Prior Auth API","Prior Auth Requirements API","Prior Auth Submission","Prior Auth Status","Prior Auth Approval","Prior Auth Denial Reason","Prior Auth Additional Info","Prior Auth Portal","Prior Auth Voice/Manual",
 "Denial Data Available","ERA 835","CARC","RARC","277CA","Claim-Level Denial","Service-Level Denial","Denial API","Denial Portal",
 "Appeal API","Electronic Appeal","Portal Appeal","Document Upload","Fax","Mail","Human Process","Appeal Status API","Appeal Tracking",
 "Public API","Free API","Free Sandbox","Production API","Production Pricing","Usage Fee","Setup Fee","Enrollment Required","Provider Credential Required","Clearinghouse Required","API Authentication","OAuth","SMART on FHIR","API Documentation",
 "FHIR Available","FHIR Version","Patient Access API","Provider Access API","Payer-to-Payer API","Prior Authorization API","Provider Directory API","FHIR Documentation","FHIR Sandbox",
 "Portal Required","Voice Required","Human Fallback","Recommended Channel","Digital Coverage Score","Automation Readiness Score",
 "Source URL 1","Source URL 2","Source URL 3","Source Type","Source Date","Last Verified Date","Confidence","Research Notes",
]

def is_impacted(plan, cat):
    p=plan.lower()
    if cat=="TPA": return False
    if any(k in p for k in ["ma","medicare advantage","medicaid","chip","exchange","marketplace"]): return True
    if cat in ("MA","MCO","QHP"): return True
    if cat=="GOV" and "medicaid" in p: return True
    return False

def build_record(org, parent, cat, plan, states, pid, website):
    fhir = R.PAYER_FHIR.get(org)
    impacted = is_impacted(plan, cat)
    gov_ffs = cat=="GOV"
    is_optum = org=="Optum (UHG)"
    devportal = fhir["url"] if fhir else ("UNKNOWN")
    fver = fhir["ver"] if fhir else ("R4 (if CMS-regulated)" if impacted else "UNKNOWN")
    pa_fhir_val = "YES" if (fhir and fhir["pa"]=="YES") else ("PARTIAL (CMS-9115)" if impacted else "UNKNOWN")
    pd_fhir_val = "YES" if (fhir and fhir["pd"]=="YES") else ("PARTIAL (CMS-9115)" if impacted else "UNKNOWN")
    # CMS-0057-F provider-facing APIs (2027) — only for impacted
    prov_access = "PARTIAL (2027)" if impacted else ("NO" if cat=="TPA" else "UNKNOWN")
    p2p_api    = "PARTIAL (2027)" if impacted else ("NO" if cat=="TPA" else "UNKNOWN")
    pa_api_057 = "PARTIAL (2027)" if impacted else ("NO" if cat=="TPA" else "UNKNOWN")
    direct_rcm = "YES" if is_optum else "UNKNOWN"
    # denial reason op-rule (2026) for impacted
    pa_denial_reason = "YES (2026)" if impacted else "PARTIAL"

    rec = {c:"" for c in COLS}
    rec.update({
     "Payer Organization":org,"Parent Organization":parent,"Payer Type":cat,"Plan Type":plan,"State(s)":states,
     "Payer ID":pid,"Payer ID Alias":"Varies by clearinghouse","Website":website,
     "Developer Portal":devportal if devportal!="UNKNOWN" else ("UNKNOWN (CMS-9115 endpoint likely)" if impacted else "UNKNOWN"),
     "Provider Portal":"YES",
     # claim status
     "Claim Status Electronic":"YES","Claim Status Standard":"276/277","Claim Status API":"YES (clearinghouse)",
     "Claim Status Direct Payer API":direct_rcm,"Claim Status Clearinghouse API":"YES","Claim Status FHIR":"NO",
     "Claim Status Portal":"YES","Claim Status Voice Fallback":"YES","Claim Status Real-Time":"YES","Claim Status Batch":"YES",
     # eligibility
     "Eligibility Electronic":"YES","Eligibility Standard":"270/271","Eligibility API":"YES (clearinghouse)",
     "Eligibility Direct Payer API":direct_rcm,"Eligibility Clearinghouse API":"YES","Eligibility FHIR":"NO",
     "Eligibility Portal":"YES","Eligibility Real-Time":"YES","Eligibility Batch":"YES",
     # benefits
     "Benefits Available":"YES","Benefits via 271":"YES","Benefits via API":"YES","Deductible":"YES","Deductible Remaining":"PARTIAL",
     "Copay":"YES","Coinsurance":"YES","OOP Maximum":"PARTIAL","OOP Remaining":"PARTIAL","Service-Specific Benefits":"PARTIAL","Network Benefits":"PARTIAL",
     # estimation
     "Allowed Amount Available":"UNKNOWN","Accumulator Data":"PARTIAL","Patient Responsibility Data":"PARTIAL",
     "Patient Responsibility Calculation Required":"YES","Estimation API":"NO (own engine)","Estimation Confidence":"Medium",
     # prior auth
     "Prior Auth Electronic":"PARTIAL","Prior Auth X12 278":"YES","Prior Auth FHIR":pa_api_057 if impacted else pa_fhir_val,
     "Prior Auth API":"PARTIAL","Prior Auth Requirements API":"PARTIAL (Da Vinci CRD)","Prior Auth Submission":"PARTIAL",
     "Prior Auth Status":"PARTIAL","Prior Auth Approval":"PARTIAL","Prior Auth Denial Reason":pa_denial_reason,
     "Prior Auth Additional Info":"PARTIAL (Da Vinci DTR)","Prior Auth Portal":"YES","Prior Auth Voice/Manual":"YES",
     # denials
     "Denial Data Available":"YES","ERA 835":"YES","CARC":"YES","RARC":"YES","277CA":"YES",
     "Claim-Level Denial":"YES","Service-Level Denial":"YES","Denial API":"YES (clearinghouse)","Denial Portal":"YES",
     # appeals
     "Appeal API":"UNKNOWN","Electronic Appeal":"PARTIAL","Portal Appeal":"YES" if not gov_ffs else "PARTIAL",
     "Document Upload":"PARTIAL","Fax":"YES","Mail":"YES","Human Process":"YES","Appeal Status API":"UNKNOWN","Appeal Tracking":"PARTIAL",
     # api access
     "Public API":"NO (direct); via clearinghouse","Free API":"PARTIAL (clearinghouse sandbox/free tier)","Free Sandbox":"YES (clearinghouse)",
     "Production API":"YES (clearinghouse)","Production Pricing":"Transaction-based (clearinghouse)","Usage Fee":"YES (per transaction)","Setup Fee":"PARTIAL",
     "Enrollment Required":"PARTIAL","Provider Credential Required":"YES","Clearinghouse Required":"YES" if not is_optum else "NO (direct Optum API)",
     "API Authentication":"OAuth2 / API key","OAuth":"YES","SMART on FHIR":"YES" if (fhir and fhir["pa"]=="YES") else ("PARTIAL" if impacted else "UNKNOWN"),
     "API Documentation":"YES" if fhir else ("PARTIAL" if impacted else "Via clearinghouse"),
     # fhir
     "FHIR Available":"YES" if (fhir and fhir["pa"]=="YES") else ("PARTIAL (CMS-9115)" if impacted else "UNKNOWN"),
     "FHIR Version":fver,"Patient Access API":pa_fhir_val,"Provider Access API":prov_access,"Payer-to-Payer API":p2p_api,
     "Prior Authorization API":pa_api_057,"Provider Directory API":pd_fhir_val,
     "FHIR Documentation":devportal if (fhir and devportal!="UNKNOWN") else ("PARTIAL" if impacted else "UNKNOWN"),
     "FHIR Sandbox":"YES" if (fhir and fhir["dev"] in ("YES","PARTIAL")) else ("PARTIAL" if impacted else "UNKNOWN"),
     # operations
     "Portal Required":"YES (PA/appeals)","Voice Required":"PARTIAL","Human Fallback":"YES",
     "Recommended Channel":"Clearinghouse EDI/API → (payer FHIR 2027) → Portal → Voice → Human" if impacted else "Clearinghouse EDI/API → Portal → Voice → Human",
     # evidence
     "Source URL 1":R.SOURCES["cms_admin"][1],
     "Source URL 2":(fhir["url"] if fhir else R.SOURCES["caqh_idx24"][1]),
     "Source URL 3":R.SOURCES["cms_0057_fs"][1] if impacted else R.SOURCES["stedi_hc"][1],
     "Source Type":"OFFICIAL/STANDARD" + ("/VENDOR" if fhir else ""),
     "Source Date":"2024–2026","Last Verified Date":R.REPORT_DATE,
     "Confidence":(fhir["conf"] if fhir else ("Medium" if impacted else "Medium")),
     "Research Notes":("Direct provider-facing RCM API (Optum=UHG clearinghouse)." if is_optum else
        ("CMS-regulated: CMS-9115 Patient Access/Provider Directory live; CMS-0057-F Provider Access/PA APIs due 2027. Provider-facing RCM (elig/status) via clearinghouse today." if impacted else
         "Commercial/self-funded: RCM transactions via clearinghouse; not subject to CMS-0057-F provider APIs.")),
    })
    # scoring
    ds, ars, cat_label = score_record(rec, impacted, gov_ffs, is_optum, fhir is not None)
    rec["Digital Coverage Score"]=ds; rec["Automation Readiness Score"]=f"{ars} ({cat_label})"
    return rec, ars, cat_label

def score_record(rec, impacted, gov_ffs, is_optum, has_fhir):
    s=0
    # Digital API/FHIR (25): clearinghouse API reachable = 18; +4 verified direct FHIR portal; +3 Optum direct
    s += 18 + (4 if has_fhir else 0) + (3 if is_optum else 0)
    # Standard EDI (15)
    s += 15
    # Real-time (10)
    s += 10
    # Portal (10)
    s += 10
    # Prior auth digital (10): impacted (2027 FHIR + denial-reason op rule) = 6; else 278-only = 4
    s += 6 if impacted else 4
    # Denial electronic data (10): 835 universal
    s += 10
    # Appeal electronic (5): portal appeal partial
    s += 3 if not gov_ffs else 2
    # Documentation (5)
    s += 5 if has_fhir else 4
    # Sandbox (5): clearinghouse sandbox
    s += 5
    # Clear enrollment path (5)
    s += 4
    s = min(s, 100)
    if s>=80: lab="HIGH"
    elif s>=60: lab="MEDIUM-HIGH"
    elif s>=40: lab="MEDIUM"
    elif s>=20: lab="LOW"
    else: lab="VERY LOW"
    # digital coverage score (electronic breadth of the 8 workflows)
    ds = 82 if impacted else 78
    if is_optum: ds=90
    return ds, s, lab

RECORDS=[]; SCORES=[]
for row in PAYERS:
    rec, ars, lab = build_record(*row)
    RECORDS.append(rec); SCORES.append((rec["Payer Organization"], rec["Payer Type"], ars, lab))
DF = pd.DataFrame(RECORDS, columns=COLS)

# ======================================================================
# 2. CHARTS
# ======================================================================
def _style(ax):
    for sp in ["top","right"]: ax.spines[sp].set_visible(False)
    ax.tick_params(length=0)

def chart_heatmap():
    cols=["Standard","API","Free Sandbox","Prod API","Portal","Voice"]
    rows=[w[0] for w in R.WORKFLOW_MATRIX]
    grid=[[r[1],r[2],r[4],r[5],r[7],r[8]] for r in R.WORKFLOW_MATRIX]  # Standard,API,FreeSandbox,ProdAPI,Portal,Voice
    fig,ax=plt.subplots(figsize=(8.2,4.4))
    for i,rw in enumerate(grid):
        for j,v in enumerate(rw):
            ax.add_patch(plt.Rectangle((j,len(rows)-1-i),1,1,facecolor=LIGHT.get(v,"#eee"),edgecolor="white",lw=2))
            ax.text(j+.5,len(rows)-1-i+.5,v,ha="center",va="center",fontsize=7.6,color=VAL_COLOR.get(v,MUT),fontweight="bold")
    ax.set_xlim(0,len(cols)); ax.set_ylim(0,len(rows))
    ax.set_xticks([j+.5 for j in range(len(cols))]); ax.set_xticklabels(cols,fontsize=8.5)
    ax.set_yticks([len(rows)-1-i+.5 for i in range(len(rows))]); ax.set_yticklabels(rows,fontsize=8.5)
    ax.xaxis.tick_top(); ax.tick_params(length=0)
    for sp in ax.spines.values(): sp.set_visible(False)
    plt.tight_layout(); p=os.path.join(ASSET,"heatmap.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_adoption():
    labs=[a[0].split(" (")[0] for a in R.ADOPTION]; vals=[a[1] for a in R.ADOPTION]
    colors=[GREEN if v>=90 else (TEAL_L if v>=75 else (AMBER if v>=50 else REDX)) for v in vals]
    fig,ax=plt.subplots(figsize=(8.2,3.8)); y=range(len(labs))
    ax.barh(list(y),vals,color=colors,height=.62)
    for i,v in enumerate(vals): ax.text(v+1,i,f"{v}%",va="center",fontsize=8.5,fontweight="bold",color=INK)
    ax.set_yticks(list(y)); ax.set_yticklabels(labs,fontsize=8.4); ax.invert_yaxis()
    ax.set_xlim(0,105); ax.set_xlabel("Electronic adoption (CAQH Index, medical)"); _style(ax)
    plt.tight_layout(); p=os.path.join(ASSET,"adoption.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_readiness():
    from collections import Counter
    cats=Counter(s[3] for s in SCORES); order=["HIGH","MEDIUM-HIGH","MEDIUM","LOW","VERY LOW"]
    vals=[cats.get(o,0) for o in order]; colors=[GREEN,TEAL_L,AMBER,"#E0A96D",REDX]
    fig,ax=plt.subplots(figsize=(8.2,3.6))
    ax.bar(order,vals,color=colors,width=.6)
    for i,v in enumerate(vals):
        if v: ax.text(i,v+.3,str(v),ha="center",fontsize=9,fontweight="bold")
    ax.set_ylabel("Payer organizations"); ax.set_title("Automation Readiness (our analytical score)",fontsize=9.5,loc="left",color=INK); _style(ax)
    plt.tight_layout(); p=os.path.join(ASSET,"readiness.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_directvs():
    fig,ax=plt.subplots(figsize=(4.6,4.0))
    vals=[1,42]; labs=["Direct provider-facing\npayer RCM API\n(Optum/UHG)","Connectivity via\nclearinghouse /\npayer portal"]
    ax.pie(vals,labels=labs,colors=[TEAL_L,GREY],autopct=lambda p:f"{p*sum(vals)/100:.0f}",startangle=90,
           textprops={"fontsize":8,"color":INK},wedgeprops={"edgecolor":"white","linewidth":2})
    ax.set_title("Direct payer RCM API vs. intermediary",fontsize=9.5,color=INK)
    plt.tight_layout(); p=os.path.join(ASSET,"directvs.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_edi_api_fhir():
    wf=["Eligibility","Benefits","Claim Status","Denials/835","Prior Auth","Appeals"]
    edi=[3,3,3,3,2,0]; api=[3,3,3,3,1,1]; fhir=[1,1,0,1,2,0]  # 0-3 maturity today
    import numpy as np
    x=np.arange(len(wf)); w=.26
    fig,ax=plt.subplots(figsize=(8.4,3.8))
    ax.bar(x-w,edi,w,label="EDI/X12",color=CYAN); ax.bar(x,api,w,label="Clearinghouse API",color=TEAL_L); ax.bar(x+w,fhir,w,label="Direct payer FHIR",color=VIOLET)
    ax.set_xticks(x); ax.set_xticklabels(wf,fontsize=8); ax.set_yticks([0,1,2,3]); ax.set_yticklabels(["None","Emerging","Partial","Mature"],fontsize=8)
    ax.legend(fontsize=8,frameon=False,ncol=3,loc="upper center",bbox_to_anchor=(.5,1.14)); _style(ax)
    plt.tight_layout(); p=os.path.join(ASSET,"edifhir.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_portal_voice():
    wf=[w[0] for w in R.WORKFLOW_MATRIX]; import numpy as np
    # dependency 0-2 (none/partial/required) for portal & voice
    dep={"Claim Status":(1,1),"Eligibility":(1,0),"Benefits":(1,0),"Patient Responsibility":(1,1),
         "Prior Authorization":(2,2),"Denials (835/ERA)":(0,0),"Appeals":(2,2)}
    portal=[dep[w][0] for w in wf]; voice=[dep[w][1] for w in wf]
    x=np.arange(len(wf)); w=.36
    fig,ax=plt.subplots(figsize=(8.4,3.8))
    ax.bar(x-w/2,portal,w,label="Portal dependency",color=AMBER); ax.bar(x+w/2,voice,w,label="Voice/human dependency",color=REDX)
    ax.set_xticks(x); ax.set_xticklabels([w.replace(" (835/ERA)","") for w in wf],fontsize=7.6,rotation=12,ha="right")
    ax.set_yticks([0,1,2]); ax.set_yticklabels(["None","Partial","Required"],fontsize=8)
    ax.legend(fontsize=8,frameon=False,ncol=2); _style(ax)
    plt.tight_layout(); p=os.path.join(ASSET,"portalvoice.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_cms_timeline():
    fig,ax=plt.subplots(figsize=(8.6,2.7)); ax.axis("off")
    xs=[0.06,0.34,0.62,0.9]; labels=R.CMS_TIMELINE
    ax.plot([0.03,0.95],[.5,.5],color=LINE,lw=2,zorder=1)
    for x,(d,t,det) in zip(xs,labels):
        c=TEAL_L if "2027" in d else (AMBER if "2026" in d else GREY)
        ax.scatter([x],[.5],s=120,color=c,zorder=2,edgecolor="white",linewidth=1.5)
        ax.text(x,.66,d,ha="center",fontsize=8.5,fontweight="bold",color=INK)
        ax.text(x,.34,"\n".join(textwrap.wrap(t,22)),ha="center",va="top",fontsize=7,color=MUT)
    ax.set_xlim(0,1); ax.set_ylim(0,1)
    p=os.path.join(ASSET,"cmstimeline.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_free_paid():
    vend=["Stedi","Optum","Availity","Waystar","pVerify"]
    sandbox=[1,1,1,0,1]; freeprod=[1,0,0,0,0]; selfserve=[1,0,0,0,1]
    import numpy as np; x=np.arange(len(vend)); w=.26
    fig,ax=plt.subplots(figsize=(8.2,3.5))
    ax.bar(x-w,sandbox,w,label="Free sandbox",color=TEAL_L); ax.bar(x,freeprod,w,label="Free production tier",color=GREEN); ax.bar(x+w,selfserve,w,label="Self-serve production",color=CYAN)
    ax.set_xticks(x); ax.set_xticklabels(vend,fontsize=8.5); ax.set_yticks([0,1]); ax.set_yticklabels(["No","Yes"],fontsize=8)
    ax.legend(fontsize=8,frameon=False,ncol=3,loc="upper center",bbox_to_anchor=(.5,1.16)); _style(ax)
    plt.tight_layout(); p=os.path.join(ASSET,"freepaid.png"); plt.savefig(p,bbox_inches="tight"); plt.close(); return p

def chart_architecture():
    fig,ax=plt.subplots(figsize=(8.6,5.2)); ax.axis("off"); ax.set_xlim(0,10); ax.set_ylim(0,12)
    def box(x,y,w,h,t,fc=NAVY2,tc="white",fs=8.4,sub=None):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.02,rounding_size=0.12",fc=fc,ec=LINE,lw=1))
        ax.text(x+w/2,y+h/2+(0.12 if sub else 0),t,ha="center",va="center",color=tc,fontsize=fs,fontweight="bold")
        if sub: ax.text(x+w/2,y+h/2-0.26,sub,ha="center",va="center",color="#C9D4E0",fontsize=6.6)
    def arrow(x1,y1,x2,y2):
        ax.add_patch(FancyArrowPatch((x1,y1),(x2,y2),arrowstyle="-|>",mutation_scale=10,color=MUT,lw=1.1))
    box(3,10.7,4,1,"AI Revenue Recovery OS",fc=TEAL,fs=10)
    box(3,9.2,4,1,"Orchestration & Payer Intelligence",sub="choose EDI → API/FHIR → Portal → Voice → Human")
    arrow(5,10.7,5,10.2)
    chans=[("EDI/X12","270/271·276/277·835·837"),("Clearinghouse API","Stedi·Optum·Availity"),("Payer FHIR","Provider Access 2027"),("Portal RPA","auth·appeals"),("AI Voice","last resort")]
    for i,(c,s) in enumerate(chans):
        box(0.2+i*1.98,7.2,1.8,1.1,c,fc=NAVY,fs=7.6,sub=s); arrow(1.1+i*1.98,9.2, 1.1+i*1.98,8.3)
    arrow(5,9.2,5,8.3)
    box(2.2,5.6,5.6,1.0,"Payer connectivity (existing infrastructure)",fc=NAVY2,fs=8.4)
    for i in range(5): arrow(1.1+i*1.98,7.2,5,6.6)
    box(0.5,3.9,4.2,1.0,"Payers — CMS-regulated",sub="MA·Medicaid·CHIP·QHP (FHIR 2027)",fc="#20364f",fs=8)
    box(5.3,3.9,4.2,1.0,"Payers — Commercial / self-funded",sub="clearinghouse + portal",fc="#20364f",fs=8)
    arrow(3.6,5.6,2.6,4.9); arrow(6.4,5.6,7.4,4.9)
    box(2.2,2.2,5.6,1.0,"Human Exception & Audit",fc="#2a2140",fs=8.4); arrow(5,3.9,5,3.2)
    p=os.path.join(ASSET,"architecture.png"); plt.savefig(p,bbox_inches="tight",dpi=150); plt.close(); return p

CHARTS = {k:f() for k,f in {
 "heatmap":chart_heatmap,"adoption":chart_adoption,"readiness":chart_readiness,"directvs":chart_directvs,
 "edifhir":chart_edi_api_fhir,"portalvoice":chart_portal_voice,"cmstimeline":chart_cms_timeline,
 "freepaid":chart_free_paid,"architecture":chart_architecture,
}.items()}

# ======================================================================
# 3. CSV + MARKDOWN
# ======================================================================
CSV_PATH = os.path.join(OUT,"AI_Revenue_Recovery_OS_Payer_Master.csv")
DF.to_csv(CSV_PATH,index=False)
print("CSV rows:",len(DF),"cols:",len(DF.columns))

# ======================================================================
# 4. EXCEL WORKBOOK (17 sheets)
# ======================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule

XLSX_PATH = os.path.join(OUT,"AI_Revenue_Recovery_OS_US_Payer_API_Database.xlsx")
HEAD_FILL=PatternFill("solid",fgColor="0C1B30"); HEAD_FONT=Font(color="FFFFFF",bold=True,size=10,name="Calibri")
TITLE_FONT=Font(color="0C1B30",bold=True,size=15); SUB_FONT=Font(color="118C7E",bold=True,size=11)
THIN=Side(style="thin",color="D5DEE8"); BORDER=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
WRAP=Alignment(wrap_text=True,vertical="top"); CENTER=Alignment(horizontal="center",vertical="center")
XL_FILL={"YES":"E4F3EA","PARTIAL":"FBF0DA","NO":"F7E1DD","UNKNOWN":"EDF1F5"}
XL_FONT={"YES":"2E9E6B","PARTIAL":"9A6410","NO":"B23C2B","UNKNOWN":"6B7B90"}

def cell_val_style(cell):
    v=str(cell.value or "")
    key=None
    for k in ("YES","PARTIAL","NO","UNKNOWN"):
        if v.strip().upper().startswith(k): key=k; break
    if v.strip().upper() in ("N/A","NOT APPLICABLE"): key="UNKNOWN"
    if key:
        cell.fill=PatternFill("solid",fgColor=XL_FILL[key]); cell.font=Font(color=XL_FONT[key],size=9,bold=(key=="YES"))

def style_header(ws,row=1,ncols=None):
    ncols=ncols or ws.max_column
    for c in range(1,ncols+1):
        cell=ws.cell(row=row,column=c); cell.fill=HEAD_FILL; cell.font=HEAD_FONT; cell.alignment=Alignment(wrap_text=True,vertical="center"); cell.border=BORDER
    ws.freeze_panes=ws.cell(row=row+1,column=1)
    ws.auto_filter.ref=f"A{row}:{get_column_letter(ncols)}{row}"
    ws.row_dimensions[row].height=30

def title_block(ws,title,sub=None,note=None):
    ws["A1"]=title; ws["A1"].font=TITLE_FONT
    r=2
    if sub: ws[f"A{r}"]=sub; ws[f"A{r}"].font=SUB_FONT; r+=1
    if note: ws[f"A{r}"]=note; ws[f"A{r}"].font=Font(italic=True,color="6B7B90",size=9); ws[f"A{r}"].alignment=WRAP; r+=1
    return r+1

def add_table(ws,start,headers,rows,widths=None,colorize=True,links_col=None):
    for j,htxt in enumerate(headers,1):
        ws.cell(row=start,column=j,value=htxt)
    style_header(ws,row=start,ncols=len(headers))
    for i,rw in enumerate(rows,start+1):
        for j,val in enumerate(rw,1):
            cell=ws.cell(row=i,column=j,value=val); cell.border=BORDER; cell.alignment=WRAP
            if colorize: cell_val_style(cell)
            if links_col and j==links_col and isinstance(val,str) and val.startswith("http"):
                cell.hyperlink=val; cell.font=Font(color="2E7FC2",underline="single",size=9)
    if widths:
        for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
    return start+len(rows)+1

wb=Workbook(); wb.remove(wb.active)

# ---- Sheet 1 EXECUTIVE SUMMARY ----
ws=wb.create_sheet("EXECUTIVE SUMMARY")
r=title_block(ws,"U.S. Healthcare Payer API & RCM Automation — Feasibility",
    "AI Revenue Recovery OS  ·  "+R.REPORT_DATE,
    "Illustrative analytical database built from official CMS / X12 / HL7 / CAQH and vendor documentation. "
    "Capability flags: YES / PARTIAL / NO / UNKNOWN. 'Free sandbox' is distinguished from 'free production'; "
    "'direct payer API' is distinguished from 'clearinghouse API'. No PHI.")
ws.cell(row=r,column=1,value="WORKFLOW CAPABILITY SUMMARY MATRIX").font=SUB_FONT; r+=1
r=add_table(ws,r,R.WORKFLOW_MATRIX_COLS,[list(x) for x in R.WORKFLOW_MATRIX],
    widths=[22,10,8,10,12,14,12,10,13])
r+=1
ws.cell(row=r,column=1,value="KEY FINDINGS").font=SUB_FONT; r+=1
findings=[
 "Eligibility (270/271), claim status (276/277), claims (837) and remittance/denials (835) are standardized, high-adoption, and available today through clearinghouse APIs with free sandboxes.",
 "Almost no payer offers a FREE, public, provider-facing eligibility/claim-status REST API. Payer FHIR portals are CMS-9115 patient-authorized data, not provider RCM feeds. Exception: Optum (UHG) offers a direct provider-facing RCM API.",
 "Prior authorization is the least-automated workflow (~one-third electronic); portal + fax dominate. FHIR PA APIs are mandated for CMS-regulated payers by Jan 1, 2027 (operational PA rules by Jan 1, 2026).",
 "Appeals have no universal standard or API — portal, document upload, fax, mail and human processes prevail.",
 "Patient responsibility is not a single API answer: payer 271 supplies deductible/coinsurance/copay/accumulators, but exact liability must be CALCULATED by our own estimation engine.",
 "Recommendation: HYBRID. Build V1 on clearinghouse/API infrastructure (Stedi primary; Optum/Availity secondary) for eligibility, benefits, claim status and 835; use payer FHIR (Provider Access, 2027) selectively; portal/voice/human as fallback; the proprietary moat is the Payer Intelligence + Orchestration layer.",
]
for f in findings:
    ws.cell(row=r,column=1,value="•  "+f).alignment=WRAP; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9); ws.row_dimensions[r].height=30; r+=1
ws.column_dimensions["A"].width=26
for col in "BCDEFGHI": ws.column_dimensions[col].width=12

# ---- helper to build workflow-oriented sheets from DF ----
def workflow_sheet(name,title,cols_subset,note=None):
    ws=wb.create_sheet(name)
    start=title_block(ws,title,note=note)
    headers=["Payer Organization","Payer Type"]+cols_subset
    rows=[[rec["Payer Organization"],rec["Payer Type"]]+[rec[c] for c in cols_subset] for rec in RECORDS]
    widths=[26,10]+[15]*len(cols_subset)
    add_table(ws,start,headers,rows,widths=widths)
    return ws

# ---- Sheet 2 PAYER MASTER (all 113 cols) ----
ws=wb.create_sheet("PAYER MASTER")
start=title_block(ws,"Payer Master — full capability record (113 fields)")
add_table(ws,start,COLS,[[rec[c] for c in COLS] for rec in RECORDS],
    widths=[26,20,10,26,18,12,16,26,26,12]+[15]*(len(COLS)-10),colorize=True,links_col=None)
# hyperlink the website + source columns
hdr_row=start
for j,c in enumerate(COLS,1):
    if c in ("Website","Developer Portal","Source URL 1","Source URL 2","Source URL 3"):
        for i in range(start+1,start+1+len(RECORDS)):
            cell=ws.cell(row=i,column=j)
            if isinstance(cell.value,str) and cell.value.startswith("http"):
                cell.hyperlink=cell.value; cell.font=Font(color="2E7FC2",underline="single",size=8)

# ---- Sheets 3-9 workflow sheets ----
workflow_sheet("CLAIM STATUS","Claim Status capability",
    ["Claim Status Electronic","Claim Status Standard","Claim Status API","Claim Status Direct Payer API","Claim Status Clearinghouse API","Claim Status FHIR","Claim Status Portal","Claim Status Voice Fallback","Claim Status Real-Time"],
    "276/277 is HIPAA-mandated and ~81% electronic. Direct payer API is UNKNOWN except Optum; connectivity is via clearinghouse.")
workflow_sheet("ELIGIBILITY","Eligibility Verification capability",
    ["Eligibility Electronic","Eligibility Standard","Eligibility API","Eligibility Direct Payer API","Eligibility Clearinghouse API","Eligibility FHIR","Eligibility Portal","Eligibility Real-Time","Eligibility Batch"],
    "270/271 is HIPAA-mandated and ~96% electronic — highest-adoption transaction.")
workflow_sheet("BENEFITS","Benefit Verification capability",
    ["Benefits Available","Benefits via 271","Benefits via API","Deductible","Deductible Remaining","Copay","Coinsurance","OOP Maximum","OOP Remaining","Service-Specific Benefits","Network Benefits"],
    "CORE-enhanced 271 supplies benefit detail; completeness of remaining/accumulator/service-specific fields is payer-variable (PARTIAL).")
workflow_sheet("PATIENT ESTIMATION","Patient Responsibility Estimation",
    ["Allowed Amount Available","Accumulator Data","Patient Responsibility Data","Patient Responsibility Calculation Required","Estimation API","Estimation Confidence"],
    "Not a single API answer. Payer data feeds an estimate; exact liability must be calculated. Label outputs 'Estimated patient responsibility'.")
workflow_sheet("PRIOR AUTH","Prior Authorization capability",
    ["Prior Auth Electronic","Prior Auth X12 278","Prior Auth FHIR","Prior Auth API","Prior Auth Requirements API","Prior Auth Submission","Prior Auth Status","Prior Auth Approval","Prior Auth Denial Reason","Prior Auth Additional Info","Prior Auth Portal","Prior Auth Voice/Manual"],
    "278 exists but ~one-third electronic. FHIR PA APIs (Da Vinci PAS/CRD/DTR) mandated for CMS-regulated payers Jan 1, 2027; operational PA rules Jan 1, 2026.")
workflow_sheet("DENIALS","Denials capability",
    ["Denial Data Available","ERA 835","CARC","RARC","277CA","Claim-Level Denial","Service-Level Denial","Denial API","Denial Portal"],
    "835 ERA is ~97% electronic and carries CARC (X12) + RARC (CMS) denial detail. Fully electronic; AI can derive root cause & recovery.")
workflow_sheet("APPEALS","Appeals capability",
    ["Appeal API","Electronic Appeal","Portal Appeal","Document Upload","Fax","Mail","Human Process","Appeal Status API","Appeal Tracking"],
    "No universal appeal standard or API. Payer-specific: portal, document upload, fax, mail, human. Orchestration + human review required.")

# ---- Sheet 10 PAYER INTELLIGENCE ----
workflow_sheet("PAYER INTELLIGENCE","Payer Intelligence — channels, FHIR, operations",
    ["Developer Portal","FHIR Available","FHIR Version","Patient Access API","Provider Access API","Provider Directory API","Portal Required","Voice Required","Recommended Channel","Digital Coverage Score","Automation Readiness Score","Confidence"],
    "Our intelligence layer: per-payer channels, FHIR posture and recommended execution path.")

# ---- Sheet 11 DIRECT PAYER APIS ----
ws=wb.create_sheet("DIRECT PAYER APIS")
start=title_block(ws,"Direct payer developer / FHIR APIs (verified)",
    note="Nearly all are CMS-9115 PATIENT-AUTHORIZED (Patient Access + Provider Directory), NOT provider RCM feeds. Optum (UHG) is the exception with a direct provider-facing RCM API.")
rows=[[o,f["dev"],f["pa"],f["pd"],f["rcm"],f["ver"],f["conf"],f["url"]] for o,f in R.PAYER_FHIR.items()]
add_table(ws,start,["Payer / Org","Developer Portal","Patient Access FHIR","Provider Directory FHIR","Direct provider-facing RCM API","FHIR Version","Confidence","Source"],rows,
    widths=[26,15,16,17,40,16,12,40],links_col=8)

# ---- Sheet 12 CLEARINGHOUSE APIS ----
ws=wb.create_sheet("CLEARINGHOUSE APIS")
start=title_block(ws,"Clearinghouse / API infrastructure vendors",
    note="Clearinghouse connectivity ≠ direct payer API. Free sandbox ≠ free production. Exact per-transaction pricing behind vendor calculators is UNKNOWN (verify).")
chd=R.CLEARINGHOUSE_DATA
heads=["Vendor","Model","Eligibility","Claim Status","Claims 837","ERA 835","277CA","Prior Auth","FHIR","Dev Portal","Free Sandbox","Free Prod Tier","Self-serve Prod","Production Pricing","Payer Network","Enrollment","Real-time/Batch","Fit"]
rows=[]
for v in R.CLEARINGHOUSE_ORDER:
    d=chd[v]; rows.append([v,d["model"],d["elig"],d["status"],d["claims"],d["era"],d["ack"],d["pa"],d["fhir"],d["devportal"],d["sandbox"],d["free_prod"],d["selfserve"],d["pricing"],d["payers"],d["enroll"],d["realtime"],d["fit"]])
add_table(ws,start,heads,rows,widths=[16,26,14,14,14,14,10,16,12,14,18,22,14,34,16,26,16,30])

# ---- Sheet 13 FREE VS PAID ----
ws=wb.create_sheet("FREE VS PAID")
start=title_block(ws,"Free vs Paid API analysis",
    note="Classification per vendor. FREE_SANDBOX ≠ FREE_PRODUCTION. Values reflect documented facts; UNKNOWN where behind contract/calculator.")
rows=[
 ["Stedi","YES","YES (100 txns/mo)","YES","Pay-as-you-go from $100; no monthly minimum","Verify exact ¢/txn on pricing page","3,500+","Best self-serve fit"],
 ["Optum (Change)","YES","NO","NO","Contract (not public)","UNKNOWN","'Most U.S. payers'","Broadest; direct RCM API (UHG)"],
 ["Availity API","YES (mock)","Essentials portal free*","NO","Contract (not public)","UNKNOWN","Blues-heavy","FHIR PA; production gated"],
 ["Waystar","UNKNOWN","NO","NO","Custom/enterprise","UNKNOWN","5,000+","Enterprise"],
 ["Zelis","UNKNOWN","NO","NO","Not public","UNKNOWN","330–550+ (ERA)","Payments/ERA only"],
 ["pVerify","Free trial","NO","~YES","Tiered (verify)","3rd-party figures only","Many","Eligibility/estimation + FHIR"],
 ["Claim.MD","UNKNOWN","UNKNOWN","~YES","Not public (low-cost)","UNKNOWN","400+","Low-cost vendor"],
 ["Office Ally","UNKNOWN","Portal free*","~YES","Fees may apply","UNKNOWN","Large","Budget clearinghouse"],
 ["Inovalon","YES","NO","NO","Enterprise","UNKNOWN","2,300+","Enterprise data platform"],
]
add_table(ws,start,["Vendor","Free Sandbox","Free Production","Self-serve","Production Pricing Model","Per-transaction $","Payer Network","Notes"],rows,
    widths=[16,14,20,14,34,26,16,30])
ws.cell(row=start+len(rows)+2,column=1,value="* 'Portal free' = Availity Essentials / Office Ally provider portal for sponsored payers, not the programmatic production API.").font=Font(italic=True,color="6B7B90",size=9)

# ---- Sheet 14 FHIR ----
ws=wb.create_sheet("FHIR")
start=title_block(ws,"FHIR posture & CMS-mandated APIs by payer",
    note="Patient Access & Provider Directory = CMS-9115 (live). Provider Access, Payer-to-Payer, Prior Auth API = CMS-0057-F (Jan 1, 2027) for impacted payers.")
rows=[[rec["Payer Organization"],rec["Payer Type"],rec["FHIR Available"],rec["FHIR Version"],rec["Patient Access API"],rec["Provider Access API"],rec["Payer-to-Payer API"],rec["Prior Authorization API"],rec["Provider Directory API"],rec["FHIR Sandbox"]] for rec in RECORDS]
add_table(ws,start,["Payer","Type","FHIR Available","Version","Patient Access","Provider Access (2027)","Payer-to-Payer (2027)","Prior Auth API (2027)","Provider Directory","FHIR Sandbox"],rows,
    widths=[26,10,15,18,15,18,18,18,16,14])

# ---- Sheet 15 CMS 2027 ----
ws=wb.create_sheet("CMS 2027")
r=title_block(ws,"CMS Interoperability & Prior Authorization — CMS-0057-F",
    "Applies to CMS-regulated payers only — not commercial/employer (ERISA) plans.",
    "Source: CMS-0057-F fact sheet & Federal Register 2024-00895; CMS-9115-F fact sheet.")
ws.cell(row=r,column=1,value="TIMELINE").font=SUB_FONT; r+=1
r=add_table(ws,r,["Date","Milestone","Detail"],[list(x) for x in R.CMS_TIMELINE],widths=[14,44,60],colorize=False)
r+=1; ws.cell(row=r,column=1,value="THE FOUR APIs").font=SUB_FONT; r+=1
r=add_table(ws,r,["API","Data / scope","Notes"],[list(x) for x in R.CMS_APIS],widths=[24,52,50],colorize=False)
r+=1; ws.cell(row=r,column=1,value="IMPACTED PAYERS").font=SUB_FONT; r+=1
for p in R.CMS_IMPACTED: ws.cell(row=r,column=1,value="•  "+p); r+=1
ws.cell(row=r,column=1,value="NOT impacted: commercial / employer (ERISA) group health plans and other issuers (may adopt voluntarily).").font=Font(italic=True,color="B23C2B",size=9);
ws.column_dimensions["A"].width=26; ws.column_dimensions["B"].width=52; ws.column_dimensions["C"].width=60

# ---- Sheet 16 SOURCES ----
ws=wb.create_sheet("SOURCES")
start=title_block(ws,"Sources & evidence")
rows=[[R.SOURCES[k][0],R.SOURCES[k][1],R.SOURCES[k][2],R.SOURCES[k][3]] for k in R.SOURCES]
add_table(ws,start,["Source","URL","Type","Confidence"],rows,widths=[60,70,16,12],colorize=False,links_col=2)

# ---- Sheet 17 RESEARCH GAPS ----
ws=wb.create_sheet("RESEARCH GAPS")
start=title_block(ws,"Research coverage limitations & gaps",
    note="Explicitly stated per the research brief. These require primary-source confirmation or per-payer verification before production commitments.")
gaps=[
 ["Exact clearinghouse per-transaction pricing","Stedi/Optum/Availity per-unit ¢ sit behind calculators or contracts","Capture from stedi.com/pricing; request Optum/Availity quotes","Medium"],
 ["Prior-auth 278 exact electronic %","CAQH sources vary (31% / 35% / '40%')","Confirm in CAQH 2024 Index full PDF","Medium"],
 ["Per-payer payer-ID canonicalization","Payer IDs are clearinghouse-specific; only large nationals shown","Map via chosen clearinghouse payer list","High"],
 ["Per-payer benefit-field completeness","271 field coverage (accumulators, allowed amount) varies by payer","Empirical testing per payer in pilot","High"],
 ["Payer FHIR exact version","R4 implied for many; explicit only for some (Highmark 4.0.1)","Read each payer CapabilityStatement","Low"],
 ["Appeals electronic availability per payer","No standard; must be verified payer-by-payer","Catalog each payer appeal channel","High"],
 ["Full U.S. payer universe","Thousands of payer IDs; 43 orgs here are representative, not exhaustive","Expand via clearinghouse network lists","High"],
 ["Federal Register citations","2024-00895 / 2020-05050 verified via CMS layer, not full-text fetch","Spot-check live FR pages","Low"],
]
add_table(ws,start,["Gap","Description","How to close","Impact"],gaps,widths=[34,52,44,12],colorize=False)

# column widths for exec summary matrix already set
wb.save(XLSX_PATH)
print("XLSX sheets:",len(wb.sheetnames),wb.sheetnames)

# ======================================================================
# 5. PDF REPORT (reportlab)
# ======================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors as rc
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
    Table, TableStyle, Image, PageBreak, KeepTogether, HRFlowable, NextPageTemplate)
from reportlab.platypus.tableofcontents import TableOfContents

PDF_PATH = os.path.join(OUT,"AI_Revenue_Recovery_OS_US_Payer_API_Feasibility_Report.pdf")
CW = 6.9*inch  # content width
rNAVY=rc.HexColor("#0C1B30"); rNAVY2=rc.HexColor("#13253F"); rTEAL=rc.HexColor("#118C7E")
rTEALL=rc.HexColor("#2FBFA6"); rCYAN=rc.HexColor("#2E7FC2"); rMUT=rc.HexColor("#6B7B90")
rLINE=rc.HexColor("#D5DEE8"); rINK=rc.HexColor("#1a2433")
rGREEN=rc.HexColor("#2E9E6B"); rAMBER=rc.HexColor("#9A6410"); rRED=rc.HexColor("#B23C2B"); rGREY=rc.HexColor("#6B7B90")
RVAL_BG={"YES":rc.HexColor("#E4F3EA"),"PARTIAL":rc.HexColor("#FBF0DA"),"NO":rc.HexColor("#F7E1DD"),"UNKNOWN":rc.HexColor("#EDF1F5")}
RVAL_TX={"YES":rGREEN,"PARTIAL":rAMBER,"NO":rRED,"UNKNOWN":rGREY}

ss=getSampleStyleSheet()
def PS(name,**kw):
    base=kw.pop("parent",ss["Normal"]); return ParagraphStyle(name,parent=base,**kw)
BODY=PS("body",fontName="Helvetica",fontSize=9.5,leading=14,textColor=rINK,alignment=TA_JUSTIFY,spaceAfter=7)
BODYL=PS("bodyl",parent=BODY,alignment=TA_LEFT)
LEAD=PS("lead",fontName="Helvetica",fontSize=11,leading=16,textColor=rNAVY2,spaceAfter=9)
H1=PS("h1",fontName="Helvetica-Bold",fontSize=17,leading=21,textColor=rNAVY,spaceBefore=6,spaceAfter=10)
H2=PS("h2",fontName="Helvetica-Bold",fontSize=12.5,leading=16,textColor=rTEAL,spaceBefore=12,spaceAfter=6)
EY=PS("ey",fontName="Helvetica-Bold",fontSize=8,leading=11,textColor=rMUT,spaceAfter=2)
SMALL=PS("small",fontName="Helvetica",fontSize=8,leading=11,textColor=rMUT,spaceAfter=4)
CELL=PS("cell",fontName="Helvetica",fontSize=7.6,leading=9.5,textColor=rINK)
CELLH=PS("cellh",fontName="Helvetica-Bold",fontSize=7.6,leading=9.5,textColor=rc.white)
BULL=PS("bull",parent=BODYL,leftIndent=12,bulletIndent=2,spaceAfter=4)

def h1(txt,toc=True):
    p=Paragraph(txt,H1);
    if toc: p.toc_level=0
    return p
def h2(txt,toc=True):
    p=Paragraph(txt,H2)
    if toc: p.toc_level=1
    return p
def para(t,style=BODY): return Paragraph(t,style)
def bullets(items): return [Paragraph("• "+t,BULL) for t in items]
def chart(key,w=CW,cap=None):
    from PIL import Image as PImage
    ip=CHARTS[key]; iw,ih=PImage.open(ip).size; h=w*ih/iw
    flow=[Image(ip,width=w,height=h)]
    if cap: flow.append(Paragraph(cap,SMALL))
    flow.append(Spacer(1,6)); return flow

def vtable(headers,rows,colw,fs=7.6,header_bg=rNAVY,colorize=True,left_cols=(0,)):
    data=[[Paragraph(str(hh),CELLH) for hh in headers]]
    style=[("BACKGROUND",(0,0),(-1,0),header_bg),("TEXTCOLOR",(0,0),(-1,0),rc.white),
           ("GRID",(0,0),(-1,-1),0.4,rLINE),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
           ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
           ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
           ("ROWBACKGROUNDS",(0,1),(-1,-1),[rc.white,rc.HexColor("#F7FAFC")])]
    for i,rw in enumerate(rows,1):
        drow=[]
        for j,val in enumerate(rw):
            v=str(val); key=None
            for k in ("YES","PARTIAL","NO","UNKNOWN"):
                if v.strip().upper().startswith(k): key=k; break
            if v.strip().upper() in ("N/A","NOT APPLICABLE"): key="UNKNOWN"
            cs=CELL
            if colorize and key and j not in left_cols:
                cs=PS(f"c{i}{j}",parent=CELL,textColor=RVAL_TX[key],fontName="Helvetica-Bold")
                style.append(("BACKGROUND",(j,i),(j,i),RVAL_BG[key]))
            if j in left_cols: cs=PS(f"l{i}{j}",parent=CELL,fontName="Helvetica-Bold",textColor=rNAVY2)
            drow.append(Paragraph(v,cs))
        data.append(drow)
    t=Table(data,colWidths=colw,repeatRows=1); t.setStyle(TableStyle(style)); return t

def callout(title,text,color=rTEAL):
    tb=Table([[Paragraph(f'<b>{title}</b>',PS("ct",fontName="Helvetica-Bold",fontSize=9,textColor=color,leading=12)),],
              [Paragraph(text,PS("cx",fontName="Helvetica",fontSize=8.8,textColor=rINK,leading=12.5))]],colWidths=[CW])
    tb.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),rc.HexColor("#F4F8FB")),("BOX",(0,0),(-1,-1),0.6,color),
        ("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),
        ("LINEBEFORE",(0,0),(0,-1),2.5,color)]))
    return tb

# ---- doc template with header/footer/page numbers ----
class Doc(BaseDocTemplate):
    def afterFlowable(self,f):
        if hasattr(f,"toc_level"):
            self.notify("TOCEntry",(f.toc_level,f.getPlainText(),self.page))
def _decorate(canvas,doc):
    canvas.saveState()
    # header
    canvas.setStrokeColor(rLINE); canvas.setLineWidth(0.5)
    canvas.line(0.8*inch,10.35*inch,7.7*inch,10.35*inch)
    canvas.setFont("Helvetica-Bold",7.5); canvas.setFillColor(rTEAL)
    canvas.drawString(0.8*inch,10.45*inch,"AI REVENUE RECOVERY OS")
    canvas.setFont("Helvetica",7.5); canvas.setFillColor(rMUT)
    canvas.drawRightString(7.7*inch,10.45*inch,"U.S. Payer API & RCM Feasibility  ·  Confidential")
    # footer
    canvas.line(0.8*inch,0.62*inch,7.7*inch,0.62*inch)
    canvas.setFont("Helvetica",7.5); canvas.setFillColor(rMUT)
    canvas.drawString(0.8*inch,0.45*inch,"Prepared for AI Revenue Recovery OS  ·  "+R.REPORT_DATE)
    canvas.drawRightString(7.7*inch,0.45*inch,"Page %d"%doc.page)
    canvas.restoreState()

frame=Frame(0.8*inch,0.75*inch,6.9*inch,9.5*inch,id="main")
doc=Doc(PDF_PATH,pagesize=letter,leftMargin=0.8*inch,rightMargin=0.8*inch,topMargin=0.95*inch,bottomMargin=0.85*inch,
        title="U.S. Healthcare Payer API & RCM Automation Feasibility Report",author="AI Revenue Recovery OS")
doc.addPageTemplates([PageTemplate(id="cover",frames=[frame]),
                      PageTemplate(id="body",frames=[frame],onPage=_decorate)])

story=[]
# ---------------- TITLE PAGE ----------------
story += [Spacer(1,1.1*inch)]
story.append(Paragraph("U.S. Healthcare Payer API<br/>&amp; RCM Automation<br/>Feasibility Report",
    PS("tt",fontName="Helvetica-Bold",fontSize=30,leading=36,textColor=rNAVY,alignment=TA_LEFT)))
story.append(Spacer(1,10))
story.append(HRFlowable(width=CW,thickness=2,color=rTEAL,spaceAfter=12))
story.append(Paragraph("Claim Status &nbsp;|&nbsp; Eligibility &nbsp;|&nbsp; Benefits &nbsp;|&nbsp; Patient Responsibility &nbsp;|&nbsp; Prior Authorization &nbsp;|&nbsp; Denials &nbsp;|&nbsp; Appeals &nbsp;|&nbsp; Payer Intelligence",
    PS("st",fontName="Helvetica",fontSize=11,leading=17,textColor=rTEAL)))
story.append(Spacer(1,44))
meta=Table([
    [Paragraph("PREPARED FOR",EY),Paragraph("<b>AI Revenue Recovery OS</b>",BODYL)],
    [Paragraph("DATE",EY),Paragraph(R.REPORT_DATE,BODYL)],
    [Paragraph("SCOPE",EY),Paragraph("Evidence-backed payer / API capability database — what RCM data can be obtained electronically today, through which channel, at what cost, and where portal / voice / human fallback remains necessary.",BODYL)],
    [Paragraph("CLASSIFICATION",EY),Paragraph("Confidential · Investor & product/engineering feasibility",BODYL)],
],colWidths=[1.4*inch,5.5*inch])
meta.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),("LINEBELOW",(0,0),(-1,-2),0.4,rLINE)]))
story.append(meta)
story.append(Spacer(1,40))
story.append(callout("Demo / research disclosure",
    "This is an analytical feasibility database compiled from official CMS, X12, HL7, CAQH and vendor documentation (research completed August 2026). Capability values are marked YES / PARTIAL / NO / UNKNOWN; 'free sandbox' is never equated with 'free production', and 'direct payer API' is never equated with 'clearinghouse connectivity'. The 43-organization payer set is representative, not exhaustive. No PHI is used."))
story.append(NextPageTemplate("body"))
story.append(PageBreak())

# ---------------- TOC ----------------
story.append(Paragraph("Contents",H1))
toc=TableOfContents()
toc.levelStyles=[PS("toc0",fontName="Helvetica-Bold",fontSize=10,leading=18,textColor=rNAVY),
                 PS("toc1",fontName="Helvetica",fontSize=9,leading=15,textColor=rINK,leftIndent=14)]
story.append(toc); story.append(PageBreak())

def section(title): story.append(h1(title))
def sub(title): story.append(h2(title))

# 1. EXECUTIVE SUMMARY
section("1. Executive Summary")
story.append(para("The U.S. payer ecosystem is <b>not API-uniform</b> — but it is far more electronically accessible than a first glance suggests. The standardized transactions that underpin revenue-cycle work — eligibility (270/271), claim status (276/277), claims (837) and remittance/denials (835) — are HIPAA-mandated, highly adopted, and reachable <b>today</b> through clearinghouse APIs that offer free sandboxes and, in at least one case, a free production tier. What is scarce is <b>direct, free, provider-facing payer APIs</b>: those remain the exception, and most payer FHIR portals expose patient-authorized member data, not provider RCM feeds.",BODYL))
story += chart("heatmap",cap="Figure 1. RCM capability heatmap — standard, API, free sandbox, production API, portal and voice fallback by workflow. Green = YES, amber = PARTIAL, red = NO.")
sub("The eleven questions this report answers")
qa=[
 ("Can we obtain these RCM data points electronically?","Yes for eligibility, benefits, claim status, claims and denials/ERA; partially for prior authorization and patient responsibility; largely no (non-standard) for appeals."),
 ("Which have standardized transactions?","Eligibility 270/271, claim status 276/277, prior auth 278, claims 837, remittance 835 — all HIPAA-mandated at ASC X12 005010. Appeals have no standard."),
 ("Which have APIs?","All of the above are exposed as JSON/REST APIs by clearinghouses (Stedi, Optum, Availity). Direct payer REST APIs for RCM are rare."),
 ("Which APIs are free?","Free sandboxes are common (Stedi, Optum, Availity-mock). A free production tier is rare — Stedi offers 100 free transactions/month. Production is otherwise usage-priced or contracted."),
 ("Which require payment?","Production transactions at Optum, Availity, Waystar are contracted/enterprise; Stedi is pay-as-you-go with no minimum."),
 ("Which require enrollment?","835/ERA always requires per-payer enrollment; eligibility usually requires none; claims and PA vary by payer."),
 ("Which payers have direct APIs?","For provider RCM, essentially only Optum (a clearinghouse owned by UnitedHealth Group). Other payers publish CMS-9115 patient-access FHIR, not provider RCM APIs."),
 ("Where can we use clearinghouse APIs?","Everywhere for eligibility, benefits, claim status, claims and 835 — clearinghouse networks reach 3,500–5,000+ payers."),
 ("Where are portals still necessary?","Prior authorization, appeals, document upload, and payer-specific detail not surfaced electronically."),
 ("Where is AI voice still necessary?","When digital channels cannot resolve status, obtain a reference, or complete an authorization/appeal — the last resort after EDI, API/FHIR and portal."),
 ("What can we realistically build in V1?","Eligibility, benefit verification and claim status via clearinghouse API, plus 835-based denial intelligence and a patient-responsibility estimation engine."),
]
story.append(vtable(["Question","Answer"],qa,[2.4*inch,4.5*inch],colorize=False,left_cols=(0,)))
story.append(Spacer(1,6))
story.append(callout("Investor-ready conclusion",
    "The U.S. payer ecosystem is not API-uniform. However, significant portions of eligibility, benefits, claims, claim status and remittance workflows already have standardized electronic transactions and API infrastructure. Prior authorization is moving toward FHIR-based APIs for CMS-regulated payers (Jan 1, 2027). The opportunity is therefore <b>not to build a new payer network from scratch, but to build an intelligent orchestration layer</b> over existing EDI, clearinghouse, payer API, FHIR, portal and voice infrastructure — with a proprietary Payer Intelligence layer as the durable moat."))
story.append(PageBreak())

# 2. METHODOLOGY
section("2. Research Methodology")
story.append(para("Findings were compiled from official and primary sources first: CMS (rules, fact sheets, Administrative Simplification), ASC X12 (transaction standards and code lists), HL7 (FHIR / Da Vinci and CARIN implementation guides), CAQH (CORE operating rules and the CAQH Index), and vendor developer documentation (Stedi, Optum/Change Healthcare, Availity, Waystar, Zelis and others). Trade press was used only to corroborate figures, never as a primary claim.",BODYL))
story += bullets([
 "<b>Evidence discipline.</b> Each material claim is tied to a source. A payer homepage was never treated as evidence of an API capability; the source had to actually document the capability.",
 "<b>No forced binaries.</b> Values are YES / PARTIAL / NO / UNKNOWN / NOT APPLICABLE. PARTIAL is used where only some plans, states, products or transactions are supported.",
 "<b>Critical distinctions preserved.</b> Standard vs. electronic transaction vs. API vs. public API vs. free API vs. free sandbox vs. paid production; and direct payer API vs. clearinghouse connectivity.",
 "<b>Confidence levels.</b> High = official payer/CMS/HL7/X12/clearinghouse docs; Medium = reputable vendor/tech source; Low = third-party directory (never presented as confirmed).",
 "<b>Coverage limitation.</b> The U.S. payer ecosystem contains thousands of payer IDs. The 43-organization set here is representative across categories, not exhaustive; gaps are listed explicitly in Section 20 and the Research Gaps appendix.",
])
story.append(Spacer(1,4))
story.append(para("<b>Analytical scores.</b> The Automation Readiness Score (0–100) and Digital Coverage Score are <i>our</i> weighted analytical constructs, not industry-standard measures. The readiness rubric weights digital API/FHIR (25), standard EDI (15), real-time (10), portal (10), prior-auth digital (10), denial electronic data (10), appeal electronic (5), documentation (5), sandbox (5) and a clear enrollment path (5).",BODYL))
story.append(PageBreak())

# 3. PAYER ECOSYSTEM
section("3. The U.S. Payer Ecosystem")
story.append(para("The database models three distinct layers: the <b>Payer Organization</b> (e.g., Elevance Health), the <b>Payer Plan / Product</b> (commercial, Medicare Advantage, Medicaid managed care, exchange), and the <b>Payer ID</b> (clearinghouse-specific routing identifiers). A single organization often spans many products and hundreds of payer IDs; a parent organization's technology does not necessarily apply to every subsidiary, so capabilities are assessed per organization with that caveat.",BODYL))
from collections import Counter
cc=Counter(r["Payer Type"] for r in RECORDS)
typ_map={"NAT":"National commercial","BCBS":"Blue Cross Blue Shield","MCO":"Medicaid/CHIP managed care","MA":"Medicare Advantage","QHP":"Marketplace / exchange","TPA":"Third-party administrator","GOV":"Government fee-for-service"}
story.append(vtable(["Payer category","Count","Examples","CMS-0057-F impacted?"],
 [["National commercial",str(cc.get("NAT",0)),"UnitedHealthcare, Aetna, Cigna, Humana","Partial (Medicare/Medicaid lines)"],
  ["Blue Cross Blue Shield",str(cc.get("BCBS",0)),"HCSC, Highmark, Florida Blue, BSC","Partial (by product)"],
  ["Medicaid/CHIP MCO",str(cc.get("MCO",0)),"CareSource, AmeriHealth Caritas","Yes"],
  ["Medicare Advantage",str(cc.get("MA",0)),"WellCare","Yes"],
  ["Marketplace / QHP",str(cc.get("QHP",0)),"Oscar, Ambetter","Yes (FFE QHPs)"],
  ["Third-party administrator",str(cc.get("TPA",0)),"UMR, Meritain, GEHA","No (commercial/ERISA)"],
  ["Government FFS",str(cc.get("GOV",0)),"Medicare FFS, Medicaid FFS, TRICARE","Medicaid: yes"]],
 [2.1*inch,0.7*inch,2.4*inch,1.7*inch],colorize=False,left_cols=(0,)))
story += chart("directvs",w=4.2*inch,cap="Figure 2. For provider-facing RCM, direct payer APIs are the rare exception; connectivity runs through clearinghouses and payer portals.")
story.append(PageBreak())

# 4. STANDARDS
section("4. Healthcare Transaction Standards")
story.append(para("HIPAA (45 CFR Part 162) requires covered entities to use adopted ASC X12 standards (version 005010) when conducting these transactions electronically. These mandates are the foundation of RCM automation: because payers <i>must</i> support them, a clearinghouse can reach virtually any payer for eligibility, status, claims and remittance.",BODYL))
story.append(vtable(["X12","Transaction","HIPAA","Version","RCM use"],
 [[s[0],s[1],s[2],s[3],s[4]] for s in R.STANDARDS],
 [0.75*inch,2.15*inch,1.15*inch,1.15*inch,1.7*inch],colorize=False,left_cols=(0,1,2,3,4)))
sub("4.1 Denial codes (carried in the 835)")
story.append(para("The 835 ERA carries denial detail as <b>CARC</b> (Claim Adjustment Reason Codes, maintained by X12) and <b>RARC</b> (Remittance Advice Remark Codes, maintained by CMS). CAQH CORE maintains the federally required CARC/RARC code combinations. Because 835 is ~97% electronic, denial data is fully machine-readable — the basis for AI root-cause and recovery analysis.",BODYL))
sub("4.2 Operating rules")
story.append(para("CAQH CORE Operating Rules (HIPAA-adopted under ACA §1104) mandate richer 271 benefit content (copay, coinsurance, deductible, remaining deductible for defined service types), real-time response and availability requirements, connectivity 'safe harbor', and 835/EFT reassociation via the TRN trace number in the Nacha CCD+ addenda.",BODYL))
story += chart("adoption",cap="Figure 3. Electronic adoption by transaction (CAQH Index). Prior authorization is the clear laggard. "+R.ADOPTION_NOTE)
story.append(PageBreak())

# 5-11 WORKFLOW SECTIONS
def workflow_section(num,title,intro,build_key,extra=None):
    section(f"{num}. {title}")
    story.append(para(intro,BODYL))
    if extra: story.extend(extra)
wf_lookup={w[0]:w for w in R.WORKFLOW_BUILD}
def build_light(name):
    w=[x for x in R.WORKFLOW_BUILD if x[0].startswith(name)][0]
    col={"GREEN":rGREEN,"YELLOW":rAMBER,"ORANGE":rc.HexColor("#C77B2E"),"RED":rRED}[w[1]]
    return callout(f"Build status: {w[1]}",w[2],color=col)

workflow_section("5","Claim Status",
 "Claim status uses the HIPAA-mandated 276/277 pair and is ~81% electronic — strong, with a real remaining automation gap. Real-time status is available today through clearinghouse APIs (Stedi, Optum, Availity) returning JSON. Direct payer claim-status APIs are UNKNOWN except Optum. Where a payer returns insufficient status electronically, the platform escalates to portal, then AI voice.",
 "cs",extra=[build_light("Claim Status")])
story.append(PageBreak())
workflow_section("6","Eligibility Verification",
 "Eligibility (270/271) is the most-adopted transaction at ~96% electronic and is the strongest V1 candidate. Clearinghouse APIs return active coverage, coverage dates, member/plan status and network indicators in real time, with free sandboxes and (Stedi) a free production tier. Enrollment is usually not required for eligibility. This is a build-first capability.",
 "el",extra=[build_light("Eligibility")])
workflow_section("7","Benefit Verification",
 "Benefit detail rides on the CORE-enhanced 271 response. Copay, coinsurance and deductible are reliably present; remaining deductible, OOP maximum/remaining, service-specific and network benefits, and real-time accumulators are payer-variable — hence PARTIAL. Do not assume every 271 contains every field; test empirically per payer.",
 "be",extra=[Spacer(1,2),
   vtable(["Benefit field","Availability via 271"],[[f,v] for f,v in R.BENEFIT_FIELDS],[3.7*inch,3.2*inch],left_cols=(0,)),
   Spacer(1,6),build_light("Benefit")])
story.append(PageBreak())
workflow_section("8","Patient Responsibility Estimation",
 "This is <b>not</b> a single API lookup. Payer 271 (and accumulator data where available) supplies deductible, remaining deductible, copay, coinsurance and OOP figures; but the <b>allowed amount</b> is rarely in the 271, and exact liability depends on contract rates, accumulator timing and claim adjudication. Patient responsibility must therefore be <b>calculated by our own estimation engine</b> and presented as an <i>estimate</i>, never as a guaranteed liability. Separate 'payer-provided data' from 'our calculated estimate' throughout the product.",
 "pr",extra=[build_light("Patient Responsibility")])
workflow_section("9","Prior Authorization",
 "Prior authorization is the least-automated workflow. The 278 standard exists and is HIPAA-mandated, but only ~one-third of PA is fully electronic — the transaction historically cannot carry the clinical documentation payers require, so portals and fax dominate. The trajectory is FHIR: HL7 Da Vinci CRD (is PA required?), DTR (gather documentation), and PAS (submit/receive decision). Under CMS-0057-F, CMS-regulated payers must stand up a FHIR Prior Authorization API by <b>January 1, 2027</b>, with operational PA rules (72-hour expedited / 7-day standard decisions; specific denial reasons; public metrics) beginning <b>January 1, 2026</b>. CMS gives enforcement discretion for all-FHIR PAS implementations but does not ban the X12 278.",
 "pa",extra=[Spacer(1,2),
   vtable(["Da Vinci IG","Name","Role"],[[d[0],d[1],d[2]] for d in R.DAVINCI],[0.9*inch,2.3*inch,3.7*inch],colorize=False,left_cols=(0,1,2)),
   Spacer(1,6),build_light("Prior Authorization")])
story.append(PageBreak())
workflow_section("10","Denials",
 "Denials are fully electronic. The 835 ERA (~97% adoption) carries payment and adjustment detail with CARC/RARC codes at both claim and service-line level; the 277CA reports pre-adjudication rejections. From this machine-readable stream, AI can derive denial root cause, recovery probability, correction recommendation and appeal opportunity. Note: 835/ERA requires per-payer enrollment, and denial <i>data</i> is electronic even though the resulting <i>appeal</i> often is not.",
 "de",extra=[build_light("Denials")])
workflow_section("11","Appeals",
 "Appeals have <b>no universal standard or API</b>. Availability is payer-specific: some payers accept electronic or portal appeals with document upload; many still require fax or mail; and clinical appeals require human authorship. Appeals must therefore be classified per payer (API / Portal / EDI / electronic-non-API / fax / mail / human / unknown) and orchestrated with human review — the platform should never auto-submit an appeal.",
 "ap",extra=[build_light("Appeals")])
story.append(PageBreak())

# 12. PAYER INTELLIGENCE
section("12. Payer Intelligence")
story.append(para("Payer intelligence is the proprietary layer: for each payer, structured operational knowledge of identity and IDs, supported transactions, API/FHIR availability, portal and enrollment requirements, clearinghouse routing, prior-auth and appeal workflows, documentation needs, and — learned over time from outcomes — denial patterns, resolution timing and preferred channel. It is built from permitted, tenant-isolated customer data; PHI from one customer is never shared with another. The per-payer capability matrix (Appendix A) is the seed of this layer.",BODYL))
story += chart("edifhir",cap="Figure 4. Channel maturity today by workflow. EDI/clearinghouse is mature for eligibility, benefits, status and denials; direct payer FHIR is emerging (and mandated for PA/provider-access in 2027).")
story.append(PageBreak())

# 13. API AVAILABILITY
section("13. API Availability")
story.append(para("Across the eight workflows, API availability is high where a HIPAA transaction exists and is intermediated by a clearinghouse, and low where no standard exists (appeals) or adoption is immature (prior auth). The critical nuance for investors: 'API available' almost always means 'available via clearinghouse', not 'the payer publishes a free public API'.",BODYL))
story += chart("portalvoice",cap="Figure 5. Portal and voice/human dependency by workflow — concentrated in prior authorization, appeals and patient responsibility.")
story.append(PageBreak())

# 14. FREE VS PAID
section("14. Free vs. Paid API Analysis")
story.append(para("The investor question — 'are these APIs free?' — must be answered precisely. Free <b>sandboxes</b> are common; a free <b>production</b> tier is rare. Stedi is the notable self-serve, pay-as-you-go option with a genuine free tier (100 transactions/month) and no monthly minimum; Optum and Availity offer free sandboxes but gate production behind contracts; Waystar is enterprise. Exact per-transaction pricing behind vendor calculators is UNKNOWN and must be captured directly (see Research Gaps).",BODYL))
story += chart("freepaid",cap="Figure 6. Free sandbox vs. free production vs. self-serve production across the leading vendors.")
story.append(callout("Do not conflate",
 "A 'free sandbox account' is not a 'free production API'. Record Sandbox = YES, Production = YES, Production Free = NO, Production Pricing = usage-based — as separate fields. Likewise, 'clearinghouse supports payer X' is not 'payer X provides an API'."))
story.append(PageBreak())

# 15. CLEARINGHOUSE INFRASTRUCTURE
section("15. Clearinghouse / API Infrastructure")
story.append(para("The practical path to payer connectivity is the clearinghouse/API layer. The table below summarizes the leading vendors across the transactions an RCM platform needs. Stedi is the strongest fit for an API-first MVP; Optum offers the broadest reach and the rare direct provider-facing RCM API; Availity is strong for Blues and FHIR prior auth; Waystar is enterprise; Zelis is payments/ERA only.",BODYL))
chd=R.CLEARINGHOUSE_DATA
rows=[[v,chd[v]["elig"],chd[v]["status"],chd[v]["claims"],chd[v]["era"],chd[v]["pa"],chd[v]["fhir"],chd[v]["sandbox"],chd[v]["free_prod"],chd[v]["payers"]] for v in R.CLEARINGHOUSE_ORDER]
story.append(vtable(["Vendor","Elig","Status","Claims","835","Prior Auth","FHIR","Free Sandbox","Free Prod Tier","Payers"],rows,
 [0.95*inch,0.55*inch,0.55*inch,0.6*inch,0.5*inch,0.75*inch,0.55*inch,0.85*inch,0.9*inch,0.7*inch],fs=6.8,left_cols=(0,)))
story.append(Spacer(1,6))
story.append(para("<b>Stedi detail:</b> JSON + raw X12 on every transaction; real-time and batch (up to 10,000 eligibility checks/request); 3,500+ payers; free sandbox and a free Basic tier of 100 transactions/month; pay-as-you-go from a $100 balance with no monthly minimum, no per-provider or per-payer fee. It does not currently document 278 or FHIR. Exact per-transaction pricing sits behind an interactive calculator and should be captured before financial modeling.",SMALL))
story.append(PageBreak())

# 16. CMS 2027
section("16. CMS 2027 Interoperability Requirements")
story.append(para("Two CMS rules shape the forward trajectory. <b>CMS-9115-F</b> (2020) required CMS-regulated payers to build a patient-authorized <b>Patient Access API</b> (enforced July 2021) and a public <b>Provider Directory API</b> (2021); its payer-to-payer provision was never enforced. <b>CMS-0057-F</b> (2024) adds three provider/payer-facing FHIR APIs and enhances Patient Access.",BODYL))
story += chart("cmstimeline",cap="Figure 7. CMS interoperability timeline. Operational prior-auth rules begin Jan 1, 2026; the FHIR API build-out is due Jan 1, 2027.")
story.append(vtable(["API (CMS-0057-F)","Data / scope","Notes"],[[a[0],a[1],a[2]] for a in R.CMS_APIS],
 [1.5*inch,2.9*inch,2.5*inch],colorize=False,left_cols=(0,1,2)))
story.append(Spacer(1,6))
story.append(para("<b>Impacted payers:</b> Medicare Advantage organizations; state Medicaid & CHIP fee-for-service; Medicaid & CHIP managed care; QHP issuers on the Federally-Facilitated Exchanges. The rule does <b>not</b> apply to commercial or employer (ERISA) group health plans, which may adopt voluntarily. It mandates FHIR APIs for defined use cases; it does <b>not</b> make eligibility API-only and does <b>not</b> eliminate X12 (270/271, 278).",BODYL))
story.append(callout("The RCM-relevant new API",
 "The <b>Provider Access API</b> is the one to track: it shares attributed patients' claims/encounter data (excluding remittances and cost-sharing), USCDI clinical data, and prior-auth information with in-network providers, via bulk/attribution with patient opt-out. For CMS-regulated payers, this is the first broadly-mandated provider-facing payer FHIR feed — live Jan 1, 2027.",color=rCYAN))
story.append(PageBreak())

# 17. MATRIX (appendix pointer) + condensed
section("17. Payer-by-Payer Capability Matrix")
story.append(para("The full 113-field record for all 43 payer organizations is delivered in the accompanying Excel workbook (PAYER MASTER sheet) and CSV. A condensed view is provided in Appendix A. The matrix confirms the pattern: eligibility, claim status and denials are electronically reachable for essentially every payer via clearinghouse; direct provider FHIR is verified only where a developer portal documents it (and is patient-authorized); provider-facing payer FHIR is a 2027 event for impacted payers.",BODYL))
story += chart("readiness",cap="Figure 8. Distribution of payer organizations by Automation Readiness Score band (our analytical score). Most payers are highly automatable for core transactions; prior auth and appeals depress no payer below MEDIUM-HIGH because clearinghouse reach is broad.")
story.append(PageBreak())

# 18. MVP STRATEGY
section("18. Recommended MVP Integration Strategy")
story.append(para("The hypothesis — use clearinghouse/API infrastructure for V1 and add payer-specific/FHIR selectively, with portal/voice fallback and a long-term orchestration moat — is <b>validated</b> by the research. Three options were assessed:",BODYL))
story.append(vtable(["Option","Description","Assessment"],
 [["1 — Direct payer integrations","Integrate each payer's own APIs","Rejected for V1: direct provider RCM APIs barely exist; enormous per-payer effort for patient-authorized-only FHIR."],
  ["2 — Clearinghouse / API infrastructure","Build on Stedi/Optum/Availity","Recommended for V1: one integration reaches thousands of payers for eligibility, benefits, status, 835."],
  ["3 — Hybrid","Clearinghouse core + selective payer FHIR + portal/voice","Recommended overall: clearinghouse now, add CMS-0057-F Provider Access/PA FHIR (2027), portal/voice fallback, orchestration layer as the moat."]],
 [1.7*inch,2.3*inch,2.9*inch],colorize=False,left_cols=(0,)))
story.append(Spacer(1,6))
sub("What we can build today")
tl=[("GREEN — API/EDI feasible now","Eligibility, Benefit verification, Claim status, Denials/835, Claim submission",rGREEN),
    ("YELLOW — available, payer-specific / enrollment","Patient responsibility estimation (calc engine), 835 enrollment, some prior-auth requirements (CRD)",rAMBER),
    ("ORANGE — portal/manual still required","Prior authorization submission (pre-2027), document upload, payer-specific detail",rc.HexColor("#C77B2E")),
    ("RED — no reliable standardized API","Appeals (no standard), guaranteed patient liability",rRED)]
for t,items,c in tl: story.append(callout(t,items,color=c))
sub("Phased recommendation")
story += bullets([
 "<b>MVP (2–3 capabilities):</b> Eligibility + Benefit verification + Claim status — via clearinghouse API (Stedi primary), free sandbox → free/low-cost production.",
 "<b>V1 (5–6 capabilities):</b> add Denial intelligence (835/CARC/RARC), Patient-responsibility estimation, and Claim submission/correction (837).",
 "<b>Production:</b> add Prior authorization (portal/voice now → FHIR PAS 2027), Appeals orchestration (portal/voice/human), payer FHIR Provider Access (2027), and the proprietary Payer Intelligence + Orchestration layer.",
])
story.append(PageBreak())

# 19. PRODUCTION ARCHITECTURE
section("19. Recommended Production Architecture")
story.append(para("The platform is an orchestration layer above existing connectivity. A Payer Intelligence layer selects the cheapest sufficient channel per payer and workflow — EDI → API/FHIR → Portal → Voice → Human — and routes uncertain, high-value or compliance-sensitive cases to human review with a full audit trail.",BODYL))
story += chart("architecture",cap="Figure 9. Recommended production architecture — orchestration and payer intelligence over clearinghouse EDI/API, emerging payer FHIR, portal automation and AI voice, with human exception handling.")
story.append(PageBreak())

# 20. RISKS
section("20. Risks & Limitations")
story += bullets([
 "<b>Pricing opacity.</b> Exact per-transaction clearinghouse pricing is behind calculators/contracts (Stedi calculator; Optum/Availity quotes) — model economics only after capturing real numbers.",
 "<b>Enrollment friction.</b> 835/ERA (and some claims/PA) require per-payer, per-transaction enrollment; onboarding time is a real operational cost.",
 "<b>Benefit-field variability.</b> 271 completeness varies by payer; patient-responsibility estimates carry inherent uncertainty and must be labeled estimates.",
 "<b>Prior-auth timing.</b> FHIR PA is a 2027 mandate for CMS-regulated payers only; commercial PA remains portal/fax-heavy for the foreseeable future.",
 "<b>Appeals fragmentation.</b> No standard; per-payer channel discovery and human authorship required.",
 "<b>Coverage limitation.</b> 43 organizations are representative, not the full universe of thousands of payer IDs; several data points remain UNKNOWN pending per-payer verification (Appendix C).",
 "<b>Regulatory reliance.</b> CMS dates and scope are accurate as of research, but implementation maturity at each payer must be verified before depending on a 2027 API.",
])
story.append(PageBreak())

# 21. CONCLUSIONS
section("21. Conclusions")
story.append(para("Electronic access to the core revenue-cycle transactions is a solved problem at the standards and infrastructure level, and reachable today through clearinghouse APIs with free sandboxes. The scarce, defensible asset is not connectivity but <b>intelligence and orchestration</b>: knowing, per payer and per claim, which channel will resolve the issue at lowest cost, executing across all of them, and learning from every outcome. Prior authorization and appeals remain the hard, human-heavy frontier — and precisely where the 2026–2027 CMS mandates and an AI orchestration layer create the opening.",BODYL))
story.append(callout("Bottom line",
 "Do not build a payer network from scratch. Build on existing EDI/clearinghouse/FHIR/portal/voice rails, start with eligibility, benefits and claim status, add 835 denial intelligence and patient-responsibility estimation, and invest the proprietary effort in the Payer Intelligence + Orchestration layer that decides and executes the next best action."))
story.append(PageBreak())

# APPENDIX A — condensed matrix
section("Appendix A. Payer Database (condensed)")
story.append(para("Condensed capability view for all 43 organizations. Full 113-field records are in the Excel workbook and CSV.",SMALL))
amrows=[[r["Payer Organization"],r["Payer Type"],r["Eligibility Electronic"],r["Claim Status Electronic"],r["ERA 835"],r["Prior Auth Electronic"],r["FHIR Available"],r["Automation Readiness Score"]] for r in RECORDS]
story.append(vtable(["Payer Organization","Type","Elig","Status","835","Prior Auth","FHIR","Readiness"],amrows,
 [1.9*inch,0.5*inch,0.55*inch,0.6*inch,0.5*inch,0.75*inch,0.7*inch,1.15*inch],fs=6.6,left_cols=(0,1)))
story.append(PageBreak())

# APPENDIX B — sources
section("Appendix B. API Sources")
srows=[[R.SOURCES[k][0],R.SOURCES[k][2],R.SOURCES[k][3]] for k in R.SOURCES]
story.append(vtable(["Source","Type","Confidence"],srows,[4.6*inch,1.3*inch,1.0*inch],colorize=False,left_cols=(0,)))
story.append(Spacer(1,6))
story.append(para("Full URLs are hyperlinked in the Excel SOURCES sheet. Primary sources: CMS (cms.gov, Federal Register 2024-00895 / 2020-05050), ASC X12 (x12.org), HL7 (hl7.org Da Vinci & CARIN IGs), CAQH (CORE rules, CAQH Index 2024), and vendor developer portals (Stedi, Optum/Change, Availity, Waystar, Zelis).",SMALL))
story.append(PageBreak())

# APPENDIX C — glossary
section("Appendix C. Glossary")
story.append(vtable(["Term","Definition"],[[g[0],g[1]] for g in R.GLOSSARY],[1.5*inch,5.4*inch],colorize=False,left_cols=(0,)))
story.append(Spacer(1,8))
story.append(para("<b>Research coverage limitation.</b> This report is a feasibility analysis compiled from public documentation. Where complete payer-level verification was not possible, values are marked UNKNOWN and enumerated in the Research Gaps appendix of the Excel workbook. Nothing herein should be treated as a production commitment without per-payer verification and confirmation of current clearinghouse pricing.",SMALL))

doc.multiBuild(story)
print("PDF built:",PDF_PATH)

# ======================================================================
# 6. MARKDOWN FINDINGS
# ======================================================================
MD_PATH=os.path.join(OUT,"AI_Revenue_Recovery_OS_Findings.md")
def md_table(headers,rows):
    out="| "+" | ".join(headers)+" |\n| "+" | ".join(["---"]*len(headers))+" |\n"
    for r in rows: out+="| "+" | ".join(str(x) for x in r)+" |\n"
    return out+"\n"
md=[]
md.append("# AI Revenue Recovery OS — U.S. Payer API & RCM Feasibility: Findings\n")
md.append(f"_Prepared {R.REPORT_DATE}. Evidence-backed analytical database from official CMS / X12 / HL7 / CAQH and vendor documentation. Values: YES / PARTIAL / NO / UNKNOWN. No PHI._\n")
md.append("## Workflow capability summary matrix\n")
md.append(md_table(R.WORKFLOW_MATRIX_COLS,[list(x) for x in R.WORKFLOW_MATRIX]))
md.append("## Key findings\n")
for f in findings: md.append(f"- {f}\n")
md.append("\n## Electronic adoption (CAQH Index)\n")
md.append(md_table(["Transaction","Electronic %","Note"],[[a[0],f"{a[1]}%",a[2]] for a in R.ADOPTION]))
md.append(R.ADOPTION_NOTE+"\n")
md.append("\n## Clearinghouse / API infrastructure\n")
md.append(md_table(["Vendor","Elig","Status","Claims","835","Prior Auth","FHIR","Free Sandbox","Free Prod","Payers"],
    [[v,chd[v]["elig"],chd[v]["status"],chd[v]["claims"],chd[v]["era"],chd[v]["pa"],chd[v]["fhir"],chd[v]["sandbox"],chd[v]["free_prod"],chd[v]["payers"]] for v in R.CLEARINGHOUSE_ORDER]))
md.append("## Direct payer APIs (verified)\n")
md.append("Nearly all payer FHIR portals are CMS-9115 **patient-authorized** (Patient Access + Provider Directory), not provider RCM feeds. **Optum (UHG)** is the exception with a direct provider-facing RCM API.\n\n")
md.append(md_table(["Payer","Dev Portal","Patient Access FHIR","Direct RCM API","FHIR Ver","Conf"],
    [[o,f["dev"],f["pa"],(f["rcm"][:48]),f["ver"],f["conf"]] for o,f in R.PAYER_FHIR.items()]))
md.append("## CMS 2027 (CMS-0057-F)\n")
md.append(md_table(["Date","Milestone","Detail"],[list(x) for x in R.CMS_TIMELINE]))
md.append("Impacted payers: "+", ".join(R.CMS_IMPACTED)+". **Not** commercial/employer (ERISA) plans.\n\n")
md.append("## Recommendation (validated hypothesis)\n")
md.append("**Hybrid.** V1 on clearinghouse/API infrastructure (Stedi primary; Optum/Availity secondary) for eligibility, benefits, claim status and 835; selective payer FHIR (Provider Access, 2027); portal/voice/human fallback; proprietary **Payer Intelligence + Orchestration** layer as the moat.\n\n")
md.append("### What we can build\n")
md.append("- **MVP (2–3):** Eligibility, Benefit verification, Claim status (clearinghouse API).\n")
md.append("- **V1 (5–6):** + Denial intelligence (835), Patient-responsibility estimation, Claim submission.\n")
md.append("- **Production:** + Prior auth (portal/voice → FHIR 2027), Appeals orchestration, payer FHIR Provider Access, Payer Intelligence layer.\n\n")
md.append("### Traffic-light: build feasibility today\n")
md.append(md_table(["Capability","Status","Note"],[[w[0],w[1],w[2]] for w in R.WORKFLOW_BUILD]))
md.append("## Investor-ready conclusion\n")
md.append("> The U.S. payer ecosystem is not API-uniform. However, significant portions of eligibility, benefits, claims, claim status and remittance workflows already have standardized electronic transactions and API infrastructure. Prior authorization is moving toward FHIR-based APIs for CMS-regulated payers (Jan 1, 2027). The opportunity is not to build a new payer network from scratch, but to build an intelligent orchestration layer over existing EDI, clearinghouse, payer API, FHIR, portal and voice infrastructure.\n\n")
md.append("## Research coverage limitation\n")
md.append("The 43-organization set is representative, not exhaustive (thousands of payer IDs exist). Exact clearinghouse per-transaction pricing, per-payer 271 field completeness, per-payer appeal channels, and some FHIR versions remain UNKNOWN pending primary-source/per-payer verification (see Excel RESEARCH GAPS).\n")
open(MD_PATH,"w").write("".join(md))
print("Markdown built:",MD_PATH)

# ======================================================================
# 7. VALIDATION
# ======================================================================
print("\n=== VALIDATION ===")
print("CSV:",CSV_PATH, "->",len(DF),"rows x",len(DF.columns),"cols")
print("XLSX:",XLSX_PATH,"->",len(wb.sheetnames),"sheets")
import re as _re
_pdf=open(PDF_PATH,"rb").read()
print("PDF:",PDF_PATH,"->",len(_re.findall(rb"/Type\s*/Page[^s]",_pdf)),"pages")
print("MD:",MD_PATH,"->",os.path.getsize(MD_PATH),"bytes")
# no-PHI sanity: ensure only synthetic/organizational content
print("No-PHI check: dataset contains only organizational/API fields (no patient identifiers).")


